"""`shoebox variation-template`: checklist -> fill-in workbook -> segmented variation listings.

build   Pull one subset of a set (e.g. 2026 Topps Update Base) from cards.v_checklist
        and write a workbook with a row per card, sorted by card number and split
        into segments of at most 250 cards (eBay's variation cap). Qty and Price
        are blank; the Listing sheet holds the default price, the ad rate and the
        volume-discount tiers.

upload  Validate the filled-in workbook and create one multi-variation listing
        per segment through the create-variation-listings pipeline.
"""

from __future__ import annotations

import logging
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from shoebox.clients.bigquery import BigQueryClient
from shoebox.models.ebay_listing import EbayListingResult
from shoebox.pipelines.create_variation_listing import (
    _derive_default_image_path,
    list_variation_groups,
)
from shoebox.settings import get_settings
from shoebox.transforms.variation_listing_builder import build_variation_title
from shoebox.transforms.variation_template import (
    DEFAULT_PRICE_LABEL,
    DEFAULT_PROMOTE_RATE,
    DEFAULT_TIER_QUANTITIES,
    EDITABLE_COLUMNS,
    PROMOTE_RATE_LABEL,
    TEMPLATE_COLUMNS,
    TIERS_HEADER_LABEL,
    TemplateError,
    build_template_records,
    parse_listing_settings,
    parse_template_rows,
)

logger = logging.getLogger(__name__)

CHECKLIST_SHEET = "Checklist"
LISTING_SHEET = "Listing"
SUMMARY_SHEET = "Summary"

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
_EDIT_FILL = PatternFill("solid", fgColor="FFF2CC")
_COLUMN_WIDTHS = {
    "Segment": 9,
    "Card Number": 13,
    "Player": 26,
    "Team": 22,
    "Note": 10,
    "Display Name": 34,
    "Qty": 7,
    "Price": 9,
    "Set Name": 22,
    "Subset Name": 16,
    "Subset ID": 16,
}


def _sql_string(value: str) -> str:
    """Escape a value for a single-quoted BigQuery string literal."""
    return value.replace("\\", "\\\\").replace("'", "\\'")


def fetch_set_checklist(set_name: str, bq: BigQueryClient | None = None) -> list[dict[str, Any]]:
    """Every cards.v_checklist row for set_name (case-insensitive), as dicts."""
    bq = bq or BigQueryClient()
    df = bq.run_query("variation_template.sql", {"set_name": _sql_string(set_name)})
    if df is None or df.empty:
        return []
    return df.astype(object).where(df.notna(), None).to_dict("records")


def default_template_path(set_name: str, subset_name: str) -> Path:
    slug = re.sub(r"[^a-z0-9]+", "_", f"{set_name} {subset_name}".lower()).strip("_")
    return Path(get_settings().paths.exports_dir) / "variation_templates" / f"{slug}.xlsx"


def _write_checklist_sheet(ws, records: list[dict[str, Any]]) -> None:
    ws.title = CHECKLIST_SHEET
    ws.append(TEMPLATE_COLUMNS)
    for col, name in enumerate(TEMPLATE_COLUMNS, 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _EDIT_FILL if name in EDITABLE_COLUMNS else _HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = _COLUMN_WIDTHS.get(name, 14)
    for rec in records:
        ws.append([rec[c] for c in TEMPLATE_COLUMNS])

    last = len(records) + 1
    qty_col = get_column_letter(TEMPLATE_COLUMNS.index("Qty") + 1)
    price_col = get_column_letter(TEMPLATE_COLUMNS.index("Price") + 1)
    for row in range(2, last + 1):
        ws[f"{price_col}{row}"].number_format = "$#,##0.00"

    qty_rule = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0")
    qty_rule.error = "Qty must be a whole number, 0 or more"
    price_rule = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0")
    price_rule.error = "Price must be a number, 0 or more"
    ws.add_data_validation(qty_rule)
    ws.add_data_validation(price_rule)
    qty_rule.add(f"{qty_col}2:{qty_col}{last}")
    price_rule.add(f"{price_col}2:{price_col}{last}")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TEMPLATE_COLUMNS))}{last}"


def _write_listing_sheet(ws, set_name: str, subset_name: str) -> None:
    rows: list[tuple[Any, Any, str]] = [
        ("Setting", "Value", "Notes"),
        ("Set", set_name, "For reference only"),
        ("Subset", subset_name, "For reference only"),
        (DEFAULT_PRICE_LABEL, None, "Price for any card whose Price cell is blank"),
        (PROMOTE_RATE_LABEL, DEFAULT_PROMOTE_RATE, "Promoted Listings ad rate; 0 skips promotion"),
        (None, None, ""),
        ("Volume Discounts", None, "Applied to each segment's listing. Leave % blank to skip"),
        (TIERS_HEADER_LABEL, "% Off", "Add rows below for more tiers; stop at a blank row"),
    ]
    rows += [(qty, None, "") for qty in DEFAULT_TIER_QUANTITIES]
    for row in rows:
        ws.append(list(row))

    for r in (1, 7, 8):
        for c in (1, 2, 3):
            ws.cell(row=r, column=c).font = _HEADER_FONT
    for r in [4, 5] + list(range(9, 9 + len(DEFAULT_TIER_QUANTITIES))):
        ws.cell(row=r, column=2).fill = _EDIT_FILL
    ws["B4"].number_format = "$#,##0.00"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 60


def _write_summary_sheet(ws, records: list[dict[str, Any]]) -> None:
    """Per-segment card range plus live in-stock counts, for checking progress."""
    seg_col = get_column_letter(TEMPLATE_COLUMNS.index("Segment") + 1)
    qty_col = get_column_letter(TEMPLATE_COLUMNS.index("Qty") + 1)
    segs = f"{CHECKLIST_SHEET}!${seg_col}:${seg_col}"
    qtys = f"{CHECKLIST_SHEET}!${qty_col}:${qty_col}"

    ws.append(["Segment", "First Card", "Last Card", "Cards", "In Stock", "Total Qty"])
    for c in range(1, 7):
        ws.cell(row=1, column=c).font = _HEADER_FONT
    by_segment: dict[int, list[dict[str, Any]]] = {}
    for rec in records:
        by_segment.setdefault(rec["Segment"], []).append(rec)
    for r, (seg, recs) in enumerate(sorted(by_segment.items()), start=2):
        ws.append(
            [
                seg,
                recs[0]["Card Number"],
                recs[-1]["Card Number"],
                len(recs),
                f'=COUNTIFS({segs},A{r},{qtys},">0")',
                f"=SUMIFS({qtys},{segs},A{r})",
            ]
        )
    ws.append([])
    ws.append(["Segments are fixed when the template is built; edit the Segment column to move"])
    ws.append(["cards between listings. A listing holds at most 250 cards."])
    for col, width in zip("ABCDEF", (10, 13, 13, 8, 10, 11), strict=True):
        ws.column_dimensions[col].width = width


def write_template(
    path: Path, records: list[dict[str, Any]], *, set_name: str, subset_name: str
) -> None:
    wb = openpyxl.Workbook()
    _write_checklist_sheet(wb.active, records)
    _write_listing_sheet(wb.create_sheet(LISTING_SHEET), set_name, subset_name)
    _write_summary_sheet(wb.create_sheet(SUMMARY_SHEET), records)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def build_variation_template(
    set_name: str,
    subset_name: str = "Base",
    *,
    segments: int = 2,
    out: Path | None = None,
    force: bool = False,
    bq: BigQueryClient | None = None,
) -> Path:
    """Write the fill-in template for one subset of a set and return its path."""
    out = out or default_template_path(set_name, subset_name)
    if out.exists() and not force:
        raise FileExistsError(
            f"{out} already exists; pass --force to overwrite it (any Qty/Price in it is lost)"
        )

    checklist = fetch_set_checklist(set_name, bq=bq)
    if not checklist:
        raise TemplateError([f"No cards in cards.v_checklist for set {set_name!r}; check the name"])

    wanted = subset_name.casefold()
    subset_rows = [r for r in checklist if str(r.get("subset_name") or "").casefold() == wanted]
    if not subset_rows:
        available = sorted({str(r.get("subset_name")) for r in checklist})
        raise TemplateError(
            [f"No subset {subset_name!r} in {set_name!r}. Subsets: " + ", ".join(available)]
        )
    subset_ids = {r.get("subset_id") for r in subset_rows}
    if len(subset_ids) > 1:
        logger.warning("Subset %r spans several Subset IDs: %s", subset_name, sorted(subset_ids))

    records = build_template_records(subset_rows, segments)
    first = subset_rows[0]
    write_template(
        out,
        records,
        set_name=str(first.get("set_name")),
        subset_name=str(first.get("subset_name")),
    )

    sizes = Counter(r["Segment"] for r in records)
    ranges = []
    for seg in sorted(sizes):
        seg_recs = [r for r in records if r["Segment"] == seg]
        ranges.append(
            f"  segment {seg}: {seg_recs[0]['Card Number']}-{seg_recs[-1]['Card Number']} "
            f"({len(seg_recs)} cards)"
        )
    logger.info(
        "Wrote %d cards in %d segment(s) to %s\n%s",
        len(records),
        len(sizes),
        out,
        "\n".join(ranges),
    )
    if len(sizes) > segments:
        logger.warning(
            "Asked for %d segment(s) but used %d so no listing exceeds 250 cards",
            segments,
            len(sizes),
        )
    return out


def read_template(path: Path) -> tuple[list[dict[str, Any]], list[tuple[Any, ...]]]:
    """(Checklist rows as dicts, Listing sheet rows as tuples) from a template."""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    for sheet in (CHECKLIST_SHEET, LISTING_SHEET):
        if sheet not in wb.sheetnames:
            raise TemplateError([f"{path} has no {sheet!r} sheet; is it a variation template?"])

    ws = wb[CHECKLIST_SHEET]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows, ())]
    missing = [c for c in TEMPLATE_COLUMNS if c not in headers]
    if missing:
        raise TemplateError([f"{CHECKLIST_SHEET} sheet is missing column(s): {', '.join(missing)}"])
    records = [dict(zip(headers, row, strict=False)) for row in rows]

    listing_rows = list(wb[LISTING_SHEET].iter_rows(values_only=True))
    return records, listing_rows


def upload_variation_template(
    path: Path,
    *,
    publish: bool = False,
    dry_run: bool = False,
    schedule: bool = False,
    include_out_of_stock: bool = False,
    images_dir: Path | None = None,
    default_image_path: Path | None = None,
) -> list[EbayListingResult]:
    """Create one variation listing per segment of a filled-in template."""
    records, listing_rows = read_template(path)

    # Check both sheets before stopping, so one pass fixes every problem.
    problems: list[str] = []
    try:
        settings = parse_listing_settings(listing_rows)
    except TemplateError as e:
        problems += [f"Listing sheet: {p}" for p in e.problems]
        settings = None
    try:
        groups = parse_template_rows(
            records,
            default_price=settings.default_price if settings else None,
            include_out_of_stock=include_out_of_stock,
        )
    except TemplateError as e:
        problems += e.problems
    if problems:
        raise TemplateError(problems)

    logger.info(
        "%s: %d listing(s); promote rate %d%%; volume discounts %s",
        path.name,
        len(groups),
        settings.promote_rate,
        settings.volume_discount_tiers or "none",
    )
    for gid, rows in groups.items():
        logger.info(
            "  segment %s: %d cards, total qty %d, $%.2f-$%.2f | %s",
            gid.rsplit("|", 1)[1],
            len(rows),
            sum(r.qty for r in rows),
            min(r.price for r in rows),
            max(r.price for r in rows),
            build_variation_title(rows),
        )

    if dry_run:
        # Image uploads write to GCS; a dry run should leave nothing behind.
        if images_dir or default_image_path:
            logger.info("Dry run: skipping image uploads")
        images_dir = default_image_path = None
    elif default_image_path is None:
        first = next(iter(groups.values()))[0]
        default_image_path = _derive_default_image_path(
            first.set_name, get_settings().paths.set_images_dir
        )

    return list_variation_groups(
        groups,
        publish=publish,
        dry_run=dry_run,
        schedule=schedule,
        images_dir=images_dir,
        default_image_path=default_image_path,
        promote_rate=settings.promote_rate,
        volume_discount_tiers=settings.volume_discount_tiers or None,
    )
