from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from shoebox.pipelines.create_variation_listing import (
    _discover_card_images,
    load_checklist_from_excel,
)
from shoebox.pipelines.variation_template import (
    build_variation_template,
    upload_variation_template,
)
from shoebox.transforms.variation_listing_builder import build_group_key
from shoebox.transforms.variation_template import (
    DISPLAY_NAME_MAX,
    MAX_VARIATIONS,
    TemplateError,
    build_template_records,
    card_sort_key,
    default_display_name,
    parse_listing_settings,
    parse_template_rows,
    segment_sizes,
)

SET = "2026 Topps Update"


def _checklist_row(n: int, subset: str = "Base", **extra) -> dict:
    row = {
        "card_id": f"{SET}|{subset}|US{n}",
        "subset_id": f"2026TU-{subset.upper()}",
        "set_year": "2026",
        "set_name": SET,
        "subset_name": subset,
        "subset_type": "Base" if subset == "Base" else "Insert",
        "derived_card_number": str(n),
        "card_number": f"US{n}",
        "player": f"Player {n}",
        "team": "Team",
        "note": None,
        "note_check": None,
    }
    row.update(extra)
    return row


def _record(n: int, segment: int = 1, qty=1, price=1.0, **extra) -> dict:
    rec = {
        "Segment": segment,
        "Card Number": f"US{n}",
        "Player": f"Player {n}",
        "Team": "Team",
        "Note": None,
        "Display Name": f"US{n} Player {n}",
        "Qty": qty,
        "Price": price,
        "Set Year": "2026",
        "Set Name": SET,
        "Subset Name": "Base",
        "Subset Type": "Base",
        "Subset ID": "2026TU-BASE",
        "Derived Card Number": str(n),
        "Record ID": f"r{n}",
    }
    rec.update(extra)
    return rec


# --- sorting / segmenting ---------------------------------------------------


def test_card_sort_key_orders_numerically_not_lexically():
    cards = ["US10", "US2", "US100", "US1"]
    assert sorted(cards, key=card_sort_key) == ["US1", "US2", "US10", "US100"]


def test_card_sort_key_prefers_derived_number_and_puts_text_last():
    assert card_sort_key("91AS-123", "123") < card_sort_key("US200", None)
    assert card_sort_key("US5") < card_sort_key("CL")


@pytest.mark.parametrize(
    ("n", "segments", "expected"),
    [
        (350, 2, [175, 175]),
        (351, 2, [176, 175]),
        (600, 2, [200, 200, 200]),  # raised so no listing exceeds 250
        (250, 1, [250]),
        (1, 2, [1]),
        (0, 2, []),
    ],
)
def test_segment_sizes(n, segments, expected):
    sizes = segment_sizes(n, segments)
    assert sizes == expected
    assert sum(sizes) == n
    assert all(s <= MAX_VARIATIONS for s in sizes)


def test_segment_sizes_rejects_zero():
    with pytest.raises(ValueError):
        segment_sizes(10, 0)


def test_default_display_name_flags_rookies_and_fits_ebay_limit():
    assert default_display_name("US1", "Shohei Ohtani") == "US1 Shohei Ohtani"
    assert default_display_name("US7", "Rookie Guy", "RC") == "US7 Rookie Guy RC"
    assert default_display_name("US8", "Inherited", None, "RC") == "US8 Inherited RC"
    long = default_display_name("US9", " / ".join(["Somebody Longname"] * 5))
    assert len(long) <= DISPLAY_NAME_MAX


def test_build_template_records_sorts_and_splits_us1_to_us350():
    checklist = [_checklist_row(n) for n in range(350, 0, -1)]  # reversed on purpose
    records = build_template_records(checklist, segments=2)

    assert [r["Card Number"] for r in records[:3]] == ["US1", "US2", "US3"]
    assert records[174]["Card Number"] == "US175" and records[174]["Segment"] == 1
    assert records[175]["Card Number"] == "US176" and records[175]["Segment"] == 2
    assert records[-1]["Card Number"] == "US350"
    assert all(r["Qty"] is None and r["Price"] is None for r in records)
    assert records[0]["Record ID"] == f"{SET}|Base|US1"


# --- Listing sheet ----------------------------------------------------------


def test_parse_listing_settings_reads_values_and_tiers():
    rows = [
        ("Setting", "Value", "Notes"),
        ("Set", SET, None),
        ("Default Price", "$0.99", None),
        ("Promote Rate (%)", 8, None),
        (None, None, None),
        ("Volume Discounts", None, None),
        ("Min Qty", "% Off", None),
        (4, 25, None),
        (2, 10, None),
        (3, None, None),  # blank % = tier skipped
        (None, None, None),
        (10, 50, None),  # below the blank row: not part of the table
    ]
    s = parse_listing_settings(rows)
    assert s.default_price == 0.99
    assert s.promote_rate == 8
    assert s.volume_discount_tiers == [{2: 10}, {4: 25}]


def test_parse_listing_settings_blank_promote_rate_disables_promotion():
    s = parse_listing_settings([("Promote Rate (%)", None)])
    assert s.promote_rate == 0
    assert s.default_price is None
    assert s.volume_discount_tiers == []


def test_parse_listing_settings_reports_every_bad_value():
    rows = [
        ("Default Price", "abc"),
        ("Promote Rate (%)", 150),
        ("Min Qty", "% Off"),
        (1, 10),
        (2, 120),
        (3, 2.5),
    ]
    with pytest.raises(TemplateError) as exc:
        parse_listing_settings(rows)
    assert len(exc.value.problems) == 5


# --- Checklist rows ---------------------------------------------------------


def test_parse_template_rows_groups_by_segment_and_keeps_out_of_stock_cards():
    records = [
        _record(1, 1, qty=2, price=1.5),
        _record(2, 1, qty=None),  # not stocked: listed as sold out
        _record(3, 2, qty=1, price=None),  # falls back to default price
        _record(4, 2, qty=0, price=None),
    ]
    groups = parse_template_rows(records, default_price=0.99)

    assert list(groups) == ["2026TU-BASE|1", "2026TU-BASE|2"]
    seg1, seg2 = groups.values()
    assert [(r.card_number, r.qty, r.price) for r in seg1] == [("US1", 2, 1.5), ("US2", 0, 1.0)]
    assert [(r.card_number, r.qty, r.price) for r in seg2] == [("US3", 1, 0.99), ("US4", 0, 0.99)]
    assert build_group_key("2026TU-BASE|1") != build_group_key("2026TU-BASE|2")


def test_parse_template_rows_out_of_stock_cards_still_need_a_price():
    records = [_record(1, qty=None, price=None)]
    with pytest.raises(TemplateError, match="needs a Price"):
        parse_template_rows(records)

    groups = parse_template_rows(records, default_price=1.0)
    (rows,) = groups.values()
    assert rows[0].qty == 0


def test_parse_template_rows_collects_all_problems():
    records = [
        _record(1, qty="two"),
        _record(2, qty=-1),
        _record(3, price=None),
        _record(4, segment=None),
        _record(5, **{"Display Name": "Same"}),
        _record(6, **{"Display Name": "Same"}),
        _record(7, **{"Display Name": "x" * (DISPLAY_NAME_MAX + 1)}),
    ]
    with pytest.raises(TemplateError) as exc:
        parse_template_rows(records)
    text = str(exc.value)
    assert len(exc.value.problems) == 6
    for card in ("US1", "US2", "US3", "US4", "US6", "US7"):
        assert f"({card})" in text


def test_parse_template_rows_rejects_segment_over_ebay_cap():
    records = [_record(n) for n in range(1, MAX_VARIATIONS + 2)]
    with pytest.raises(TemplateError, match="eBay allows 250"):
        parse_template_rows(records)


def test_parse_template_rows_with_no_cards_says_so():
    with pytest.raises(TemplateError, match="No cards in the Checklist sheet"):
        parse_template_rows([_record(1, **{"Card Number": None})])


def test_parse_template_rows_display_name_defaults_to_card_number():
    groups = parse_template_rows([_record(1, **{"Display Name": None})])
    (rows,) = groups.values()
    assert rows[0].display_name == "US1"


# --- workbook round trip ----------------------------------------------------


class _FakeBigQuery:
    def __init__(self, rows: list[dict]):
        self.rows = rows
        self.calls: list[tuple[str, dict]] = []

    def run_query(self, sql, params=None, **_):
        self.calls.append((sql, params))
        return pd.DataFrame(self.rows)


@pytest.fixture
def update_base_bq():
    rows = [_checklist_row(n) for n in range(1, 351)]
    rows += [_checklist_row(n, subset="Rookie Debut") for n in range(1, 4)]
    rows[0]["note"] = "RC"
    rows[1]["team"] = float("nan")  # pandas NaN from BigQuery
    return _FakeBigQuery(rows)


def _fill(path, qty_price: dict[str, tuple], listing: dict[str, object]) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb["Checklist"]
    headers = [c.value for c in ws[1]]
    card_col, qty_col, price_col = (headers.index(h) + 1 for h in ("Card Number", "Qty", "Price"))
    for row in range(2, ws.max_row + 1):
        card = ws.cell(row=row, column=card_col).value
        if card in qty_price:
            qty, price = qty_price[card]
            ws.cell(row=row, column=qty_col).value = qty
            ws.cell(row=row, column=price_col).value = price
    ls = wb["Listing"]
    for row in ls.iter_rows():
        if row[0].value in listing:
            row[1].value = listing[row[0].value]
    wb.save(path)


def test_build_writes_template_and_upload_dry_run_makes_two_listings(
    tmp_path, monkeypatch, update_base_bq
):
    monkeypatch.chdir(tmp_path)  # exports/ (results JSONL) lands in tmp
    out = tmp_path / "update_base.xlsx"

    path = build_variation_template(SET.lower(), "base", out=out, bq=update_base_bq)
    assert path == out
    sql, params = update_base_bq.calls[0]
    assert sql == "variation_template.sql" and params == {"set_name": SET.lower()}

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Checklist", "Listing", "Summary"]
    ws = wb["Checklist"]
    assert ws.max_row == 351  # header + US1..US350, inserts excluded
    assert ws["B2"].value == "US1" and ws["F2"].value == "US1 Player 1 RC"
    assert ws["D3"].value is None  # NaN team written as blank
    summary = [r for r in wb["Summary"].iter_rows(min_row=2, max_row=3, values_only=True)]
    assert [r[:4] for r in summary] == [(1, "US1", "US175", 175), (2, "US176", "US350", 175)]

    with pytest.raises(FileExistsError):
        build_variation_template(SET, "Base", out=out, bq=update_base_bq)

    stock = {f"US{n}": (1, None) for n in range(1, 351) if n % 5}  # skip every 5th card
    stock["US1"] = (3, 4.99)
    _fill(out, stock, {"Default Price": 0.99, "Promote Rate (%)": 5, 2: 10, 3: 15})

    results = upload_variation_template(out, dry_run=True)

    assert len(results) == 2
    titles = [r.request.data["item_group"]["title"] for r in results]
    assert titles == [
        f"{SET} #US1-US175 You Pick, Complete Your Set!",
        f"{SET} #US176-US350 You Pick, Complete Your Set!",
    ]
    first = results[0].request.data
    offers = list(first["sku_offers"].values())
    assert len(offers) == 175  # every card, stocked or not
    assert sum(o["availableQuantity"] == 0 for o in offers) == 35  # every 5th card
    prices = sorted(o["pricingSummary"]["price"]["value"] for o in offers)
    assert prices[0] == 0.99 and prices[-1] == 4.99
    assert results[0].inventory_item_group_key != results[1].inventory_item_group_key
    assert (tmp_path / "exports" / "jsonl" / "ebay_listings.jsonl").exists()


def test_upload_reports_template_problems_before_touching_ebay(
    tmp_path, monkeypatch, update_base_bq
):
    monkeypatch.chdir(tmp_path)
    out = build_variation_template(SET, "Base", out=tmp_path / "t.xlsx", bq=update_base_bq)
    _fill(out, {"US1": (1, None)}, {})  # stocked but no price and no default

    def _boom(*a, **k):
        raise AssertionError("eBay should not be called")

    monkeypatch.setattr("shoebox.pipelines.variation_template.list_variation_groups", _boom)
    with pytest.raises(TemplateError, match="needs a Price"):
        upload_variation_template(out, publish=True)

    _fill(out, {}, {"Promote Rate (%)": 150})  # now both sheets have a problem
    with pytest.raises(TemplateError) as exc:
        upload_variation_template(out, publish=True)
    problems = exc.value.problems
    assert problems[0].startswith("Listing sheet:")
    assert problems[1].startswith("Row 2 (US1): needs a Price")
    assert len(problems) == 351  # every card needs a price, stocked or not


def test_build_unknown_subset_lists_the_ones_available(tmp_path, update_base_bq):
    with pytest.raises(TemplateError, match="Subsets: Base, Rookie Debut"):
        build_variation_template(SET, "Chrome", out=tmp_path / "x.xlsx", bq=update_base_bq)


def test_template_also_reads_as_an_inventory_workbook(tmp_path, update_base_bq):
    out = build_variation_template(SET, "Base", out=tmp_path / "t.xlsx", bq=update_base_bq)
    _fill(out, {"US2": (1, 1.25)}, {})
    groups = load_checklist_from_excel(out, in_stock_only=True)
    (rows,) = groups.values()
    assert [(r.card_number, r.price) for r in rows] == [("US2", 1.25)]


# --- per-card image matching ------------------------------------------------


def test_discover_card_images_does_not_confuse_us1_with_us10(tmp_path):
    for name in ("US10_front.jpg", "US1_back.jpg", "US1_front.jpg", "US100.jpg"):
        (tmp_path / name).write_bytes(b"")
    found = _discover_card_images(tmp_path, "US1")
    assert {k: v.name for k, v in found.items()} == {
        "front": "US1_front.jpg",
        "back": "US1_back.jpg",
    }
